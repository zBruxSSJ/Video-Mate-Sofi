"""Genera Guion_Estacion2.docx: el texto exacto que dice cada integrante,
con tiempos (m:ss) y un color por persona.
"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

COLORS = {
    "Sofía": RGBColor(0x0d, 0x94, 0x88),      # turquesa oscurecido (legible en blanco)
    "Micaela": RGBColor(0xc2, 0x55, 0x0a),    # naranja oscurecido
    "Fabrizio": RGBColor(0x6d, 0x28, 0xd9),   # violeta oscurecido
    "Jerónimo": RGBColor(0x92, 0x6c, 0x05),   # dorado oscurecido
    "TODOS": RGBColor(0x1f, 0x29, 0x37),
}

SWATCH_HEX = {
    "Sofía": "2dd4bf",
    "Micaela": "fb923c",
    "Fabrizio": "a78bfa",
    "Jerónimo": "fbbf24",
    "TODOS": "334155",
}

SCENES = [
    {
        "num": 1, "titulo": "Intro", "inicio": "0:00", "fin": "0:15",
        "lineas": [
            ("Sofía", "0:00", "Caso de estudio. Estación dos, de Matemática, tercero BATIN."),
            ("Sofía", "0:03", "La función que vamos a analizar es: h de x, igual a uno sobre, x menos dos."),
            ("Sofía", "0:07", "Somos del colegio San Ignacio de Loyola School, de Paraguay."),
            ("Sofía", "0:10", "Grupo cuatro: Sofía Cabrera, Micaela Cubilla, Fabrizio Riquelme y Jerónimo Lovera."),
        ],
    },
    {
        "num": 2, "titulo": "Condiciones de continuidad", "inicio": "0:15", "fin": "0:45",
        "lineas": [
            ("Sofía", "0:15", "Antes de resolver el caso, repasemos las condiciones de continuidad en un punto x igual a a."),
            ("Sofía", "0:20", "Condición uno: h de a tiene que existir; es decir, la función debe estar definida en ese punto."),
            ("Sofía", "0:27", "Condición dos: el límite, cuando x tiende a a, de h de x, tiene que existir, con un valor finito."),
            ("Sofía", "0:35", "Condición tres: ese límite tiene que ser igual a h de a."),
            ("Sofía", "0:41", "Si una sola de estas tres condiciones falla, la función es discontinua en ese punto."),
        ],
    },
    {
        "num": 3, "titulo": "El caso", "inicio": "0:45", "fin": "1:10",
        "lineas": [
            ("Micaela", "0:45", "Vamos al caso de estudio: la función h de x, igual a uno sobre, x menos dos."),
            ("Micaela", "0:50", "Queremos analizar qué pasa en el punto x igual a dos, que marcamos acá en la recta numérica."),
            ("Micaela", "0:58", "Fíjense que, si reemplazamos x por dos, el denominador se hace cero."),
            ("Micaela", "1:04", "Entonces la pregunta es: ¿es continua h de x, en x igual a dos?"),
        ],
    },
    {
        "num": 4, "titulo": "Condición 1", "inicio": "1:10", "fin": "1:35",
        "lineas": [
            ("Micaela", "1:10", "Empecemos con la condición uno: h de a tiene que existir. Evaluamos en a igual a dos."),
            ("Micaela", "1:15", "h de dos, es igual a uno sobre, dos menos dos."),
            ("Micaela", "1:19", "Eso es uno sobre cero."),
            ("Micaela", "1:23", "Y la división por cero no existe."),
            ("Micaela", "1:27", "Entonces, h de dos no existe: falla la condición uno, porque la función no está definida en x igual a dos."),
        ],
    },
    {
        "num": 5, "titulo": "Condición 2", "inicio": "1:35", "fin": "2:15",
        "lineas": [
            ("Fabrizio", "1:35", "Pasemos a la condición dos: el límite, cuando x tiende a dos, de h de x, tiene que existir."),
            ("Fabrizio", "1:40", "Para eso, calculamos los límites laterales: por izquierda y por derecha."),
            ("Fabrizio", "1:46", "Por el lado izquierdo, probamos con valores que se acercan a dos, pero son menores: uno coma nueve, uno coma noventa y nueve, uno coma novecientos noventa y nueve."),
            ("Fabrizio", "1:55", "Y h de x nos da: menos diez, menos cien, menos mil. Es decir, el límite lateral izquierdo es menos infinito."),
            ("Fabrizio", "2:05", "Por el lado derecho, con dos coma uno, dos coma cero uno, dos coma cero cero uno, h de x nos da: diez, cien, mil. El límite lateral derecho es más infinito."),
            ("Fabrizio", "2:13", "Menos infinito es distinto de más infinito: entonces no existe el límite cuando x tiende a dos."),
        ],
    },
    {
        "num": 6, "titulo": "Condición 3 y conclusión parcial", "inicio": "2:15", "fin": "2:40",
        "lineas": [
            ("Fabrizio", "2:15", "Ahora la condición tres: el límite tiene que ser igual a h de a."),
            ("Fabrizio", "2:19", "En este caso, h de dos es distinto del límite, cuando x tiende a dos, de h de x, porque ninguno de los dos existe."),
            ("Fabrizio", "2:27", "Entonces tampoco se cumple la condición tres."),
            ("Fabrizio", "2:30", "En resumen: no se cumplen las tres condiciones de continuidad."),
        ],
    },
    {
        "num": 7, "titulo": "Gráfico", "inicio": "2:40", "fin": "3:20",
        "lineas": [
            ("Jerónimo", "2:40", "Miremos esto en el gráfico de h de x, igual a uno sobre, x menos dos."),
            ("Jerónimo", "2:46", "Acá tenemos la recta x igual a dos, que es una asíntota vertical."),
            ("Jerónimo", "2:52", "Por la izquierda, la curva baja, y baja, hasta irse a menos infinito."),
            ("Jerónimo", "2:59", "Por la derecha, la curva sube, y sube, hasta irse a más infinito."),
            ("Jerónimo", "3:06", "Las dos ramas se acercan a la asíntota, pero nunca la tocan."),
            ("Jerónimo", "3:12", "Por eso concluimos que h de x es discontinua en x igual a dos, y presenta una discontinuidad inevitable de salto infinito."),
        ],
    },
    {
        "num": 8, "titulo": "Error detectado y cierre", "inicio": "3:20", "fin": "3:50",
        "lineas": [
            ("Jerónimo", "3:20", "Por último, les contamos un error muy común."),
            ("Jerónimo", "3:24", "Muchos piensan que, si el límite 'da infinito', entonces la función es continua."),
            ("Jerónimo", "3:30", "Pero eso es incorrecto: un límite infinito, o un límite que no existe, significa justamente que la función NO es continua."),
            ("Jerónimo", "3:38", "La conclusión correcta es: h de x es discontinua en x igual a dos, con una discontinuidad inevitable de salto infinito, porque la recta x igual a dos es asíntota vertical."),
            ("TODOS", "3:48", "¡Gracias por ver!"),
        ],
    },
]


def set_cell_shading(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def build_document():
    doc = Document()

    section = doc.sections[0]
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_heading("Guion — Estación 2: Caso de Estudio", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run(
        "San Ignacio de Loyola School (Paraguay) · 3° BATIN · Matemática\n"
        "Grupo 4: Sofía Cabrera · Micaela Cubilla · Fabrizio Riquelme · Jerónimo Lovera\n"
        "Duración total del video: 3:50 min — Video_Estacion2_CasoDeEstudio.mp4"
    )
    run.italic = True
    run.font.size = Pt(10.5)

    doc.add_paragraph()

    # --- Leyenda de colores ---
    leg_title = doc.add_paragraph()
    leg_title.add_run("Reparto por escena:").bold = True
    leg = doc.add_paragraph(
        "Sofía → Escenas 1 y 2  ·  Micaela → Escenas 3 y 4  ·  "
        "Fabrizio → Escenas 5 y 6  ·  Jerónimo → Escenas 7 y 8  "
        "(el cierre “¡Gracias por ver!” lo dicen los cuatro juntos)."
    )
    leg.runs[0].font.size = Pt(10.5)

    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    hdr = table.rows[0].cells
    hdr[0].text = "Integrante"
    hdr[1].text = "Color en este guion"
    for person in ["Sofía", "Micaela", "Fabrizio", "Jerónimo"]:
        row = table.add_row().cells
        p0 = row[0].paragraphs[0]
        r0 = p0.add_run(person)
        r0.font.color.rgb = COLORS[person]
        r0.bold = True
        set_cell_shading(row[1], SWATCH_HEX[person])

    doc.add_paragraph()
    doc.add_paragraph().add_run(
        "Indicaciones de lectura: hablar con tono natural y pausado, remarcando los "
        "términos en negrita. Los números y símbolos están escritos tal como se dicen "
        "en voz alta (por ejemplo, “uno sobre cero”, “más infinito”). Cada línea indica "
        "el minuto:segundo aproximado del video en el que empieza a decirse."
    ).italic = True

    for sc in SCENES:
        doc.add_page_break() if sc["num"] > 1 else None
        h = doc.add_heading(
            f"Escena {sc['num']} — {sc['titulo']}  ({sc['inicio']}–{sc['fin']})", level=1
        )

        for persona, tiempo, texto in sc["lineas"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(8)

            r_time = p.add_run(f"[{tiempo}]  ")
            r_time.font.size = Pt(10)
            r_time.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            r_time.italic = True

            r_name = p.add_run(f"{persona}: ")
            r_name.bold = True
            r_name.font.color.rgb = COLORS[persona]
            r_name.font.size = Pt(12)

            r_text = p.add_run(texto)
            r_text.font.size = Pt(12)
            r_text.font.color.rgb = RGBColor(0x1a, 0x1a, 0x1a)

    out_path = "/home/user/Video-Mate-Sofi/video_estacion2/Guion_Estacion2.docx"
    doc.save(out_path)
    print("Guion guardado en:", out_path)


if __name__ == "__main__":
    build_document()
